from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import subprocess
import tempfile
import uuid
from pathlib import Path

from .db import audit, connect, data_dir
from .classify import classify
from .embeddings import embed_texts, model_id
from .extract import suggest_fields
from .text_quality import looks_garbled_vietnamese


MAX_UPLOAD_BYTES = 50 * 1024 * 1024
ALLOWED_EXTENSIONS = {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".txt", ".md",
                      ".png", ".jpg", ".jpeg", ".tif", ".tiff"}


class ProcessingUnavailable(RuntimeError):
    pass


def save_blob(file_obj, filename: str) -> tuple[str, str]:
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError("Unsupported file type")
    name = f"{uuid.uuid4().hex}{extension}"
    path = data_dir() / "blobs" / name
    digest = hashlib.sha256()
    total = 0
    try:
        with path.open("wb") as target:
            while True:
                piece = file_obj.read(1024 * 1024)
                if not piece:
                    break
                total += len(piece)
                if total > MAX_UPLOAD_BYTES:
                    raise ValueError("File exceeds 50 MiB upload limit")
                digest.update(piece)
                target.write(piece)
        if total == 0:
            raise ValueError("Empty file")
    except Exception:
        path.unlink(missing_ok=True)
        raise
    return name, digest.hexdigest()


def _ocr_image(image_bytes: bytes, filename: str) -> str:
    endpoint = os.getenv("DAU_OCR_URL", "").strip()
    if not endpoint:
        raise ProcessingUnavailable("Scanned page requires DAU_OCR_URL or imported OCR pages")
    from urllib.parse import urlparse
    parsed = urlparse(endpoint)
    if parsed.scheme not in {"http", "https"} or parsed.hostname not in {"127.0.0.1", "localhost", "::1"}:
        raise ProcessingUnavailable("DAU_OCR_URL must point to a local OCR service")
    import requests
    response = requests.post(endpoint, data={"mode": os.getenv("DAU_OCR_MODE", "auto")},
                             files={"file": (filename, image_bytes, "image/png")}, timeout=180)
    response.raise_for_status()
    payload = response.json()
    text = payload.get("corrected_text") or payload.get("raw_text") or payload.get("text")
    if not isinstance(text, str):
        raise ProcessingUnavailable("OCR endpoint returned no text")
    return text


def _pdf_pages(path: Path, source_name: str = "native-pdf") -> list[dict]:
    import fitz
    pages = []
    with fitz.open(path) as pdf:
        if len(pdf) > 1000:
            raise ValueError("PDF exceeds 1000 page limit")
        for index, page in enumerate(pdf, 1):
            text = page.get_text("text")
            source = source_name
            if len(text.strip()) < 20 or looks_garbled_vietnamese(text):
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
                text = _ocr_image(pix.tobytes("png"), f"page-{index}.png")
                source = "ocr"
            pages.append({"page_no": index, "text": text, "source": source})
    return pages


def _office_pages(path: Path, filename: str) -> list[dict]:
    executable = shutil.which("soffice")
    if not executable:
        raise ProcessingUnavailable("Cần LibreOffice nội bộ để đọc DOC/XLS/XLSX")
    with tempfile.TemporaryDirectory(dir=data_dir()) as directory:
        root = Path(directory)
        source = root / ("source" + Path(filename).suffix.lower())
        shutil.copyfile(path, source)
        output = root / "converted"
        output.mkdir()
        profile = (root / "profile").as_uri()
        environment = {**os.environ, "XDG_CACHE_HOME": str(root / "cache")}
        result = subprocess.run([executable, f"-env:UserInstallation={profile}", "--headless",
                                 "--convert-to", "pdf", "--outdir", str(output), str(source)],
                                capture_output=True, text=True, timeout=180, env=environment)
        converted = output / "source.pdf"
        if result.returncode != 0 or not converted.is_file():
            raise ProcessingUnavailable("Không chuyển được tệp Office sang trang PDF")
        return _pdf_pages(converted, "converted-office")


def extract_pages(blob_name: str, filename: str) -> list[dict]:
    path = data_dir() / "blobs" / blob_name
    sidecar = path.with_name(path.name + ".ocr.json")
    if sidecar.exists():
        raw = json.loads(sidecar.read_text(encoding="utf-8"))
        if not isinstance(raw, list) or not raw or len(raw) > 1000:
            raise ValueError("OCR pages must be a nonempty list of at most 1000 pages")
        pages = []
        for index, item in enumerate(raw, 1):
            if not isinstance(item, dict) or not isinstance(item.get("text"), str):
                raise ValueError("Invalid OCR page")
            pages.append({"page_no": index, "text": item["text"], "source": "existing-ocr"})
        return pages
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return _pdf_pages(path)
    if ext in {".doc", ".xls", ".xlsx"}:
        return _office_pages(path, filename)
    if ext == ".docx":
        from docx import Document
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        for table in doc.tables:
            text += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)
        return [{"page_no": 1, "text": text, "source": "native-docx"}]
    if ext in {".txt", ".md"}:
        return [{"page_no": 1, "text": path.read_text(encoding="utf-8-sig"), "source": "native-text"}]
    if ext in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
        from PIL import Image
        import io
        image = Image.open(path).convert("RGB")
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        return [{"page_no": 1, "text": _ocr_image(buffer.getvalue(), "image.png"), "source": "ocr"}]
    raise ValueError("Unsupported file type")


def chunk_pages(pages: list[dict], size: int = 950, overlap: int = 140) -> list[dict]:
    result = []
    for page in pages:
        text = re.sub(r"[ \t]+", " ", page["text"]).strip()
        start = 0
        ordinal = 0
        while start < len(text):
            end = min(start + size, len(text))
            if end < len(text):
                split = max(text.rfind("\n", start + size // 2, end),
                            text.rfind(". ", start + size // 2, end))
                if split > start:
                    end = split + 1
            piece = text[start:end].strip()
            if piece:
                result.append({"page_no": page["page_no"], "ordinal": ordinal, "text": piece})
                ordinal += 1
            if end >= len(text):
                break
            start = max(start + 1, end - overlap)
    return result


def _propose_references(db, document_id: str, pages: list[dict]) -> None:
    # A citation must name an existing document number; all machine-created edges await review.
    for page in pages:
        for match in re.finditer(r"(?:theo|căn cứ|thay thế|sửa đổi)\s+(?:văn bản|quyết định|thông tư)?\s*(?:số\s*)?([\d]+/[\wĐđ-]+)",
                                 page["text"], re.IGNORECASE):
            number = match.group(1)
            target = db.execute(
                "SELECT d.id FROM fields f JOIN versions v ON v.id=f.version_id "
                "JOIN documents d ON d.current_version_id=v.id "
                "WHERE f.field_key='document_number' AND f.value=? AND d.id!=? LIMIT 1",
                (number, document_id),
            ).fetchone()
            if not target:
                continue
            existing = db.execute(
                "SELECT 1 FROM relations WHERE source_id=? AND target_id=? AND relation_type='cites'",
                (document_id, target["id"]),
            ).fetchone()
            if not existing:
                db.execute(
                    "INSERT INTO relations(id,source_id,target_id,relation_type,evidence,page_no,status) "
                    "VALUES (?,?,?,?,?,?,'proposed')",
                    (str(uuid.uuid4()), document_id, target["id"], "cites",
                     page["text"][max(0, match.start()-50):match.end()+50], page["page_no"]),
                )


def process_version(version_id: str) -> None:
    with connect() as db:
        version = db.execute(
            "SELECT v.*,d.kind,d.classification_status FROM versions v "
            "JOIN documents d ON d.id=v.document_id WHERE v.id=?",
            (version_id,),
        ).fetchone()
        if not version:
            raise ValueError("Unknown version")
        if version["status"] == "ready":
            return
        db.execute("UPDATE versions SET status='processing',error=NULL WHERE id=?", (version_id,))
        db.execute("UPDATE jobs SET status='processing',updated_at=CURRENT_TIMESTAMP WHERE version_id=?", (version_id,))
    try:
        with connect() as db:
            reused = db.execute(
                "SELECT v.id FROM versions v WHERE v.sha256=? AND v.blob_name=? "
                "AND v.status='ready' AND v.id!=? LIMIT 1",
                (version["sha256"], version["blob_name"], version_id),
            ).fetchone()
            cached_pages = [dict(row) for row in db.execute(
                "SELECT page_no,text,source FROM pages WHERE version_id=? ORDER BY page_no",
                (reused["id"],),
            )] if reused else []
        pages = cached_pages or extract_pages(version["blob_name"], version["filename"])
        suggestion = None
        if version["classification_status"] == "proposed":
            with connect() as db:
                source = db.execute("SELECT e.archive_path,b.kind_hint FROM import_batch_entries e "
                                    "JOIN import_batches b ON b.id=e.batch_id "
                                    "WHERE e.document_id=? LIMIT 1", (version["document_id"],)).fetchone()
            suggestion = classify(version["filename"], source["archive_path"] if source else "",
                                  "\n".join(page["text"] for page in pages[:2]),
                                  source["kind_hint"] if source else "")
        fields = suggest_fields(pages, suggestion["kind"] if suggestion else version["kind"])
        chunks = chunk_pages(pages)
        embeddings = (embed_texts([item["text"] for item in chunks])
                      if chunks and os.getenv("DAU_ENABLE_SEMANTIC", "0") == "1" else None)
        with connect() as db:
            # A worker can restart after an interrupted job. Replace partial derived
            # records atomically before inserting the new OCR result.
            db.execute("DELETE FROM chunks_fts WHERE chunk_id IN "
                       "(SELECT id FROM chunks WHERE version_id=?)", (version_id,))
            db.execute("DELETE FROM chunks WHERE version_id=?", (version_id,))
            db.execute("DELETE FROM fields WHERE version_id=?", (version_id,))
            db.execute("DELETE FROM pages WHERE version_id=?", (version_id,))
            for page in pages:
                db.execute("INSERT INTO pages(version_id,page_no,text,source) VALUES (?,?,?,?)",
                           (version_id, page["page_no"], page["text"], page["source"]))
            for field in fields:
                db.execute(
                    "INSERT INTO fields(id,version_id,field_key,value,page_no,evidence) VALUES (?,?,?,?,?,?)",
                    (str(uuid.uuid4()), version_id, field["field_key"], field["value"],
                     field["page_no"], field["evidence"]),
                )
            for index, chunk in enumerate(chunks):
                chunk_id = str(uuid.uuid4())
                db.execute(
                    "INSERT INTO chunks(id,version_id,page_no,ordinal,text,embedding,model_id) "
                    "VALUES (?,?,?,?,?,?,?)",
                    (chunk_id, version_id, chunk["page_no"], chunk["ordinal"], chunk["text"],
                     embeddings[index] if embeddings else None, model_id() if embeddings else None),
                )
                db.execute("INSERT INTO chunks_fts(text,chunk_id) VALUES (?,?)", (chunk["text"], chunk_id))
            db.execute("UPDATE versions SET status='ready' WHERE id=?", (version_id,))
            db.execute("UPDATE documents SET current_version_id=? WHERE id=?", (version_id, version["document_id"]))
            if suggestion:
                db.execute("UPDATE documents SET kind=?,subtype=?,classification_confidence=?,"
                           "classification_evidence=? WHERE id=? AND classification_status='proposed'",
                           (suggestion["kind"], suggestion["subtype"], suggestion["confidence"],
                            suggestion["evidence"], version["document_id"]))
            _propose_references(db, version["document_id"], pages)
            db.execute("UPDATE jobs SET status='ready',updated_at=CURRENT_TIMESTAMP WHERE version_id=?", (version_id,))
            audit(db, version["created_by"], "version.ready", version["document_id"], version_id)
    except Exception as exc:
        with connect() as db:
            db.execute("UPDATE versions SET status='failed',error=? WHERE id=?", (str(exc)[:500], version_id))
            db.execute("UPDATE jobs SET status='failed',error=?,updated_at=CURRENT_TIMESTAMP WHERE version_id=?",
                       (str(exc)[:500], version_id))
            audit(db, version["created_by"], "version.failed", version["document_id"], str(exc)[:500])
        raise


def reindex_embeddings() -> int:
    with connect() as db:
        rows = db.execute("SELECT id,text FROM chunks ORDER BY rowid").fetchall()
    if not rows:
        return 0
    total = 0
    for start in range(0, len(rows), 32):
        batch = rows[start:start+32]
        vectors = embed_texts([row["text"] for row in batch])
        if vectors is None:
            raise ProcessingUnavailable("Set DAU_EMBED_MODEL_PATH to a local SentenceTransformer model")
        with connect() as db:
            db.executemany("UPDATE chunks SET embedding=?,model_id=? WHERE id=?",
                           [(vector, model_id(), row["id"]) for row, vector in zip(batch, vectors)])
        total += len(batch)
    return total
